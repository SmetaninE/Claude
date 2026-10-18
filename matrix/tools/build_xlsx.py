"""Builds the factor matrix workbook from data/<type>.json + data/<type>_scores.json."""
import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SP = os.environ.get('MATRIX_ROOT') or os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TYPES = [('home', '1 Главная'), ('category', '2 Категория'), ('facet', '3 Фасет'), ('product', '4 Карточка'),
         ('series', '5 Серия'), ('article', '6 Статья'), ('news', '7 Новость'), ('reviews', '8 Отзывы')]
DOMAINS = ['ozon.ru', 'wildberries.ru', 'market.yandex.ru', 'detmir.ru', 'aliexpress.ru', 'boobl-goom.ru',
           'igroray.ru', 'ural.toys', 'markertoys.ru', 'megamarket.ru', 'toys-land.ru']
PRIO_FILL = {'Обязательный': 'F4B6B6', 'Важный': 'FCE0A8', 'Желательный': 'FFF5C2', 'Низкий': 'E7E6E6', 'Дифференциатор': 'C9B8F0'}
MISS_FILL = PatternFill('solid', fgColor='FF0000')
HEAD_FILL = PatternFill('solid', fgColor='1F3864')
thin = Side(style='thin', color='BFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def cov_score(k, n):
    if k == 0 or n == 0: return 0
    if k == n: return 5
    s = k / n
    if s >= 0.75: return 4
    if s >= 0.4: return 3
    if k == 1: return 1
    return 2

def prio(total, k, strength, n):
    # п. 5.6: у 1–2 конкурентов и при этом меньшинство (≤1/3 доступных), сила 4–5
    if 1 <= k <= 2 and k / n <= 0.34 and strength >= 4: return 'Дифференциатор'
    return 'Обязательный' if total >= 12 else 'Важный' if total >= 8 else 'Желательный' if total >= 4 else 'Низкий'

def autosize(ws, maxw=60, minw=6):
    for col in ws.columns:
        L = max((max(len(l) for l in str(c.value).split('\n')) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(minw, min(maxw, L + 2))

def header(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center'); c.border = BORDER
    ws.freeze_panes = 'A2'

def build(out):
    wb = Workbook(); wb.remove(wb.active)
    src = wb.create_sheet('Источники')
    header(src, ['Тип страницы', 'Сайт', 'URL', 'Статус', 'Причина / комментарий'])
    summary = []
    for key, label in TYPES:
        p = f'{SP}/data/{key}.json'
        if not os.path.exists(p): continue
        d = json.load(open(p)); sc = json.load(open(f'{SP}/data/{key}_scores.json'))
        s = d['sima']; src.append([label, 'sima-land.ru ({myurl})', s['url'], s['status'], s.get('reason', '')])
        urls = {u['domain']: u for u in d['urls']}
        for dom in DOMAINS:
            u = urls.get(dom, {'url': None, 'status': 'нет данных', 'reason': ''})
            src.append([label, dom, u.get('url'), u['status'], u.get('reason', '')])
        live = [dm for dm in DOMAINS if urls.get(dm, {}).get('status', '').startswith(('ok', 'частично'))]
        n = len(live)
        rows = []
        for f in d['factors']:
            s2 = sc[f['name']]
            k = sum(1 for dm in live if f['marks'].get(dm) == '+')
            q = sum(1 for dm in live if f['marks'].get(dm) == '?')
            c = cov_score(k, n); tot = c + s2['strength'] + s2['intent']
            cov_txt = f'{k}/{n} ({round(100 * k / n) if n else 0}%)' + (f', ещё {q} «?»' if q else '')
            rows.append(dict(f=f, k=k, cov=cov_txt, c=c, s=s2['strength'], i=s2['intent'], tot=tot,
                             pr=prio(tot, k, s2['strength'], n), sima=s2.get('sima', '?'),
                             com='; '.join(x for x in [s2['why'], f.get('note', ''), s2.get('sima_note', '')] if x)))
        rows.sort(key=lambda r: (-r['tot'], -r['s'], r['f']['name']))
        ws = wb.create_sheet(f'{label} — матрица')
        cols = ['Фактор', 'Описание', 'Группа', 'Тип влияния'] + live + \
               ['Охват', 'Оценка охвата', 'Сила влияния', 'Соответствие интенту', 'Итоговый балл', 'Приоритет', 'sima-land ({myurl})', 'Комментарий']
        header(ws, cols)
        for j, dm in enumerate(live):
            cell = ws.cell(row=1, column=5 + j); cell.hyperlink = urls[dm]['url']
        for r in rows:
            f = r['f']
            ws.append([f['name'], f['description'], f['group'], f['impact']] + [f['marks'].get(dm, '') for dm in live] +
                      [r['cov'], r['c'], r['s'], r['i'], r['tot'], r['pr'], r['sima'], r['com']])
            ri = ws.max_row
            pcell = ws.cell(row=ri, column=len(cols) - 2)
            pcell.fill = PatternFill('solid', fgColor=PRIO_FILL[r['pr']]); pcell.font = Font(bold=True)
            scell = ws.cell(row=ri, column=len(cols) - 1)
            scell.alignment = Alignment(horizontal='center')
            if r['sima'] == '–' and (r['pr'] == 'Обязательный' or (r['pr'] == 'Дифференциатор' and r['tot'] >= 12)):
                scell.fill = MISS_FILL; scell.font = Font(bold=True, color='FFFFFF')
                ws.cell(row=ri, column=1).font = Font(bold=True, color='C00000')
            elif r['sima'] == '–' and r['pr'] in ('Важный', 'Дифференциатор'):
                scell.fill = PatternFill('solid', fgColor='F8CBAD')
            elif r['sima'] == '±':
                scell.fill = PatternFill('solid', fgColor='FFE699')
            elif r['sima'] == '+':
                scell.fill = PatternFill('solid', fgColor='C6EFCE')
            for c_ in ws[ri]: c_.border = BORDER; c_.alignment = Alignment(wrap_text=c_.column in (2, len(cols)), vertical='top', horizontal=c_.alignment.horizontal)
            for j in range(len(live)): ws.cell(row=ri, column=5 + j).alignment = Alignment(horizontal='center', vertical='top')
        ws.auto_filter.ref = ws.dimensions
        autosize(ws)
        ws.column_dimensions['B'].width = 50; ws.column_dimensions[get_column_letter(len(cols))].width = 70
        for j in range(len(live)): ws.column_dimensions[get_column_letter(5 + j)].width = max(6, len(live[j]) + 2)
        # gaps sheet
        g = wb.create_sheet(f'{label} — пробелы')
        checked = [r for r in rows if r['sima'] in '+–±']
        have = sum(1 for r in checked if r['sima'] == '+') + 0.5 * sum(1 for r in checked if r['sima'] == '±')
        pct = round(100 * have / len(checked)) if checked else None
        imp = [r for r in rows if r['pr'] in ('Обязательный', 'Важный') or (r['pr'] == 'Дифференциатор' and r['tot'] >= 8)]
        gaps = [r for r in imp if r['sima'] in ('–', '±')]
        unk = [r for r in imp if r['sima'] == '?']
        g.append([f'{label}: {s["url"]}']); g['A1'].font = Font(bold=True, size=13)
        g.append([d['intent']])
        g.append([f'Факторов: {len(rows)}; обязательных: {sum(r["pr"] == "Обязательный" for r in rows)}; важных: {sum(r["pr"] == "Важный" for r in rows)}; дифференциаторов: {sum(r["pr"] == "Дифференциатор" for r in rows)}; конкурентов с данными: {n}/11'])
        g.append([f'Покрытие факторов на sima-land: {pct}% (из {len(checked)} проверенных; «±» = 0,5)' if pct is not None else 'Покрытие: sima-land ещё не проверена (TLS reset). Ниже — чек-лист обязательных и важных факторов, которые нужно проверить на {myurl}; после загрузки HTML в html/ колонка «sima-land» заполнится и станет списком пробелов.'])
        g.append([])
        hdr = ['Фактор', 'Приоритет', 'Итоговый балл', 'sima-land', 'Группа', 'Охват', 'Что сделать / комментарий']
        g.append(hdr); hr = g.max_row
        for c in g[hr]: c.font = Font(bold=True, color='FFFFFF'); c.fill = HEAD_FILL; c.border = BORDER
        for r in gaps + unk:
            g.append([r['f']['name'], r['pr'], r['tot'], r['sima'], r['f']['group'], r['cov'], r['com']])
            ri = g.max_row
            g.cell(row=ri, column=2).fill = PatternFill('solid', fgColor=PRIO_FILL[r['pr']])
            if r['sima'] == '–' and r['pr'] in ('Обязательный', 'Дифференциатор') and r['tot'] >= 12:
                g.cell(row=ri, column=4).fill = MISS_FILL; g.cell(row=ri, column=1).font = Font(bold=True, color='C00000')
            for c_ in g[ri]: c_.border = BORDER; c_.alignment = Alignment(wrap_text=c_.column == 7, vertical='top')
        g.freeze_panes = g.cell(row=hr + 1, column=1)
        g.auto_filter.ref = f'A{hr}:G{g.max_row}'
        for col, w in zip('ABCDEFG', [42, 16, 8, 9, 14, 16, 90]): g.column_dimensions[col].width = w
        summary.append(dict(type=label, n=len(rows), must=sum(r['pr'] == 'Обязательный' for r in rows), pct=pct,
                            gaps=[r['f']['name'] for r in gaps if r['pr'] == 'Обязательный' or r['tot'] >= 12][:6],
                            gaps_imp=[r['f']['name'] for r in gaps], live=n))
    autosize(src, maxw=90); src.auto_filter.ref = src.dimensions
    wb.move_sheet('Источники', offset=len(wb.sheetnames))
    wb.save(out)
    print(json.dumps(summary, ensure_ascii=False, indent=1))

if __name__ == '__main__':
    build(sys.argv[1])
